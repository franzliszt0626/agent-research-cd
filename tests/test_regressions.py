import asyncio
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

import httpx
from docx import Document
from pydantic import ValidationError

from app.citation_mapper import (_fallback_reference, _match_citation, _normalize_reference_shape,
                                 split_citation_claims, split_sentences)
from app.document_loader import _extract_plain_entries, load_document
from app.mcp_client import LiteratureClient, MCPFailure, decode_tool_result, exception_detail
from app.mcp_server import _select_candidate, search_literature, _get_json, _openalex_result
from app.source_fetch import parse_source_html, arxiv_identifier, public_url
from app.config import validate_base_url
from app.llm import LLMServiceError, check_llm
from app.schemas import ReportSchema, ReviewPayload
from app.workflow import WorkflowService
from app.prompts import DEEP_REVIEW_SYSTEM


GB1 = "中国电力企业联合会. 中国电力行业年度发展报告2025[R]. 北京: 中国电力企业联合会, 2025."
GB2 = "Girshick R, Donahue J, Darrell T, et al. Rich feature hierarchies for accurate object detection[C]//CVPR. 2014: 580-587."
GB3 = "Ren S, He K, Girshick R, et al. Faster R-CNN: towards real-time object detection[J]. IEEE TPAMI, 2017, 39(6): 1137-1149."


class ParserTests(unittest.TestCase):
    def test_model_cannot_invent_title_doi_or_year(self):
        ref = _normalize_reference_shape({"title": "Invented title", "doi": "10.1234/invented", "year": 2099},
                                         {"raw": GB3}, 1)
        self.assertEqual(ref["title"], "Faster R-CNN: towards real-time object detection")
        self.assertEqual(ref["doi"], "")
        self.assertEqual(ref["year"], 2017)

    def test_merged_entries_and_tail(self):
        entries = _extract_plain_entries("参考文献\n" + GB1 + " " + GB2 + "\n" + GB3 + " 致谢 这段不是参考文献 2026 数据库表")
        self.assertEqual(len(entries), 3)
        self.assertEqual(entries[-1]["raw"], GB3)
        refs = [_fallback_reference(e["raw"], i) for i, e in enumerate(entries, 1)]
        self.assertEqual(refs[0]["authors"], ["中国电力企业联合会"])
        self.assertEqual(refs[2]["title"], "Faster R-CNN: towards real-time object detection")
        self.assertEqual(refs[2]["year"], 2017)
        self.assertEqual(refs[2]["authors"], ["Ren S", "He K", "Girshick R"])

    def test_numbered_continuations_keep_label(self):
        entries = _extract_plain_entries("[5] " + GB2 + "\n[9] " + GB3)
        self.assertEqual([e["label"] for e in entries], ["5", "9"])
        ref = _fallback_reference(entries[1]["raw"], 2, label="9")
        self.assertTrue(_match_citation("检测效果改善[7–9]。", ref)[0])
        self.assertFalse(_match_citation("检测效果改善[2]。", ref)[0])

    def test_docx_boundaries_and_table_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "paper.docx"
            doc = Document()
            doc.add_paragraph("模型有改进[1]。另一项有提升[2]。")
            doc.add_table(rows=1, cols=1).cell(0, 0).text = "正文表格"
            doc.add_paragraph("参 考 文 献")
            doc.add_paragraph(GB1 + " " + GB2)
            doc.add_paragraph("致 谢")
            doc.add_paragraph("感谢大家 2026")
            doc.add_table(rows=1, cols=1).cell(0, 0).text = "数据库 VARCHAR 2026"
            doc.save(path)
            loaded = load_document(path)
            self.assertEqual(len(loaded.bibliography_entries), 2)
            self.assertIn("正文表格", loaded.main_text)
            self.assertNotIn("数据库", loaded.bibliography_text)
            self.assertEqual(len(split_sentences("模型有改进[1]。另一项有提升[2]。")), 2)

    def test_author_year_is_local_and_key_is_exact(self):
        ref = _fallback_reference("Pradhan R, Lahiri A (2022) Explainable AI. IEEE", 1)
        self.assertTrue(_match_citation("Full explanations are required (Pradhan et al. 2022).", ref)[0])
        self.assertFalse(_match_citation("Pradhan disagreed in 2020 with Smith (2022).", ref)[0])
        ref["bib_key"] = "paper"
        self.assertFalse(_match_citation("A claim [CITE:paper-extra].", ref)[0])
        self.assertTrue(_match_citation("A claim [CITE:paper,other].", ref)[0])

    def test_numeric_interval_is_not_a_citation(self):
        ref = _fallback_reference(GB1, 1)
        self.assertFalse(_match_citation("归一化坐标取值范围为 [0, 1]。", ref)[0])

    def test_compound_numeric_citations_are_split_by_claim(self):
        sentence = "Du等人提出GLM预训练框架[33]，Hu等人提出LoRA高效微调方法[34]。"
        self.assertEqual(split_citation_claims(sentence), [
            "Du等人提出GLM预训练框架[33]",
            "Hu等人提出LoRA高效微调方法[34]。",
        ])
        self.assertEqual(split_citation_claims("同一结论由多项研究支持[3,4]。"), ["同一结论由多项研究支持[3,4]。"])


class TransportTests(unittest.IsolatedAsyncioTestCase):
    def test_configured_base_url_rejects_trailing_comma(self):
        with self.assertRaisesRegex(ValueError, "LLM_BASE_URL"):
            validate_base_url("https://api.llm.ustc.edu.cn,")
        self.assertEqual(validate_base_url("https://api.llm.ustc.edu.cn/v1/"), "https://api.llm.ustc.edu.cn/v1")

    def test_arxiv_source_identifier_and_citation_metadata(self):
        self.assertEqual(arxiv_identifier("https://arxiv.org/abs/2310.11511"), "2310.11511")
        html = '''<html><head><meta name="citation_title" content="Self-RAG: learning to retrieve, generate, and critique through self-reflection">
        <meta name="citation_author" content="Akari Asai"><meta name="citation_pdf_url" content="/pdf/2310.11511"></head>
        <body><blockquote class="abstract">Abstract: We propose a retrieval framework.</blockquote>
        <main>''' + "Study findings. " * 100 + "</main></body></html>"
        result = parse_source_html(html, "https://arxiv.org/abs/2310.11511", "2310.11511")
        self.assertEqual(result["source"], "arXiv")
        self.assertEqual(result["abstract"], "We propose a retrieval framework.")
        self.assertEqual(result["pdf_url"], "https://arxiv.org/pdf/2310.11511")

    def test_webpage_description_not_claimed_as_abstract(self):
        html = '<html><head><title>国家能源局发布数据</title><meta name="description" content="网站描述"></head><body><main>' + "2023 年线路长度数据。" * 80 + '</main></body></html>'
        result = parse_source_html(html, "https://www.nea.gov.cn/article", "")
        self.assertEqual(result["abstract"], "")
        self.assertIn("2023 年", result["_page_text"])

    async def test_source_urls_must_be_public(self):
        for url in ("http://127.0.0.1/private", "http://localhost/private", "file:///etc/passwd"):
            with self.assertRaises(ValueError):
                await public_url(url)

    async def test_crossref_rate_limit_retry(self):
        request = httpx.Request("GET", "https://api.crossref.org/works")
        client = AsyncMock()
        client.get.side_effect = [httpx.Response(429, headers={"retry-after": "1"}, request=request),
                                  httpx.Response(200, json={"message": {"items": []}}, request=request)]
        with patch("app.mcp_server.asyncio.sleep", new=AsyncMock()) as pause:
            result = await _get_json(client, str(request.url))
        self.assertIn("message", result)
        self.assertEqual(client.get.await_count, 2)
        pause.assert_awaited_once_with(1.0)

    async def test_exact_source_precedes_fuzzy_search(self):
        exact = {"found": True, "title": "LoRA", "arxiv_id": "2106.09685", "abstract": "Evidence"}
        with patch("app.mcp_server.fetch_source", new=AsyncMock(return_value=exact)), \
             patch("app.mcp_server._search_openalex", new=AsyncMock()) as alex:
            result = await search_literature("LoRA", arxiv_id="2106.09685")
        self.assertEqual(result["arxiv_id"], "2106.09685")
        alex.assert_not_called()

    async def test_sufficient_openalex_metadata_skips_crossref(self):
        exact = {"found": True, "title": "A study", "doi": "10.1234/test", "abstract": "Evidence"}
        with patch("app.mcp_server._search_openalex", new=AsyncMock(return_value=exact)), \
             patch("app.mcp_server._search_crossref", new=AsyncMock()) as cross, \
             patch("app.mcp_server._unpaywall_lookup", new=AsyncMock(return_value=None)):
            result = await search_literature("A study")
        self.assertTrue(result["found"])
        cross.assert_not_called()

    def test_unknown_access_is_not_closed(self):
        result = _openalex_result({"title": "A paper", "open_access": {}})
        self.assertIsNone(result["is_open_access"])
        self.assertEqual(result["access_status"], "access-undetermined")

    def test_content_wrappers_and_error_groups(self):
        record = {"found": True, "abstract": "Test abstract"}
        for value in (record, json.dumps(record), [{"type": "text", "text": json.dumps(record)}],
                      {"structuredContent": record}, {"content": [{"type": "text", "text": json.dumps(record)}]}):
            self.assertEqual(decode_tool_result(value), record)
        with self.assertRaises(MCPFailure):
            decode_tool_result({"isError": True, "content": "boom"})
        self.assertIn("ConnectionError: refused", exception_detail(ExceptionGroup("TaskGroup", [ConnectionError("refused")])))

    async def test_retry(self):
        client = LiteratureClient()
        call = AsyncMock(side_effect=[ConnectionError("refused"), ["ok"]])
        with patch("app.mcp_client.asyncio.sleep", new=AsyncMock()):
            result = await client._retry(call)
        self.assertEqual(result, ["ok"])
        self.assertEqual(call.await_count, 2)

    async def test_provider_fallback_and_timeout(self):
        cross = {"found": True, "title": "A study", "doi": "", "abstract": "Evidence", "access_status": "access-undetermined"}
        with patch("app.mcp_server._search_openalex", new=AsyncMock(side_effect=httpx.ReadTimeout("timeout"))), \
             patch("app.mcp_server._search_crossref", new=AsyncMock(return_value=cross)):
            result = await search_literature("A study")
        self.assertTrue(result["found"])
        self.assertTrue(result["provider_errors"])
        with patch("app.mcp_server._search_openalex", new=AsyncMock(side_effect=httpx.ReadTimeout("timeout"))), \
             patch("app.mcp_server._search_crossref", new=AsyncMock(return_value=None)):
            result = await search_literature("A study")
        self.assertEqual(result["status"], "network_timeout")

    def test_wrong_search_hit_is_rejected(self):
        items = [{"title": "Unrelated agriculture", "doi": "https://doi.org/10.1234/wrong"},
                 {"title": "Faster R-CNN", "doi": "https://doi.org/10.1234/right"}]
        self.assertEqual(_select_candidate(items, "Faster R-CNN", ""), items[1])
        self.assertIsNone(_select_candidate(items, "Different topic entirely", ""))
        self.assertIsNone(_select_candidate(items, "Faster R-CNN", "10.1234/missing"))
        self.assertIsNone(_select_candidate([{"title": "QA-LoRA: Quantization-Aware Low-Rank Adaptation of Large Language Models"}],
                                             "LoRA: low-rank adaptation of large language models", ""))
        self.assertIsNone(_select_candidate([{"title": "LoRA: low-rank adaptation of large language models", "publication_year": 2023,
                                               "authorships": [{"author": {"display_name": "Wrong Person"}}]}],
                                             "LoRA: low-rank adaptation of large language models", "", authors=["Hu E J"], year=2022))


class WorkflowTests(unittest.IsolatedAsyncioTestCase):
    async def test_model_preflight_aborts_before_document_or_mcp(self):
        service = WorkflowService()
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "paper.tex"
            path.write_text("Test", encoding="utf-8")
            with patch("app.workflow.check_llm", new=AsyncMock(side_effect=LLMServiceError("connection_error", "refused"))), \
                 patch("app.workflow.load_document") as loader:
                with self.assertRaises(LLMServiceError):
                    await service.run(path)
                loader.assert_not_called()

    async def run_case(self, *, metadata=None, initial=0.9, deep=0.9, tool_failure=False, llm_failure=False, invalid=False, full_available=True):
        service = WorkflowService()
        service.literature.prepare = AsyncMock(side_effect=MCPFailure("mcp_error", "refused") if tool_failure else None)
        metadata = metadata or {"found": True, "title": "Test paper", "abstract": "Evidence supports the claim.", "access_status": "open-access"}
        async def tool(name, args):
            if name == "search_literature":
                return metadata
            return {"full_text_available": full_available, "full_text": "Evidence supports the claim. " * 100 if full_available else "",
                    "full_text_status": "retrieved" if full_available else "no-open-pdf-found"}
        service.literature.call = AsyncMock(side_effect=tool)
        async def reviewer(system, user):
            if llm_failure:
                raise ConnectionError("LLM unavailable")
            is_deep = system == DEEP_REVIEW_SYSTEM
            score = deep if is_deep else initial
            return {"confidence": "invalid" if invalid else score,
                    "verdict": "supports" if score > .5 else "does_not_support", "rationale": "Evidence reviewed",
                    "supporting_passage": "Evidence supports the claim." if is_deep else "",
                    "evidence_from_abstract": "" if is_deep else "Evidence supports the claim."}
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "paper.tex"
            path.write_text(r"A test claim \cite{test}. \begin{thebibliography}{1}\bibitem{test}Smith J (2022) Test paper. Journal.\end{thebibliography}", encoding="utf-8")
            with patch("app.citation_mapper.ainvoke_json", new=AsyncMock(return_value={})), \
                 patch("app.workflow.check_llm", new=AsyncMock(return_value={"status": "ok"})), \
                 patch("app.workflow.ainvoke_json", new=AsyncMock(side_effect=reviewer)) as llm:
                report = await service.run(path)
                llm_calls = llm.await_count
        ReportSchema.model_validate(report)
        return report, service, llm_calls

    async def test_initial_pass_skips_fulltext(self):
        report, service, calls = await self.run_case()
        self.assertEqual(report["statistics"]["passed"], 1)
        self.assertEqual(service.literature.call.await_count, 1)
        self.assertEqual(calls, 1)

    async def test_stream_reports_each_executed_step_and_final_report(self):
        service = WorkflowService()
        service.literature.prepare = AsyncMock(return_value=None)
        service.literature.call = AsyncMock(return_value={
            "found": True, "title": "Test paper", "abstract": "Evidence supports the claim.",
            "access_status": "open-access",
        })
        review = {"confidence": .9, "verdict": "supports", "rationale": "Evidence reviewed",
                  "evidence_from_abstract": "Evidence supports the claim.", "supporting_passage": ""}
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "paper.tex"
            path.write_text(r"A test claim \cite{test}. \begin{thebibliography}{1}\bibitem{test}Smith J (2022) Test paper. Journal.\end{thebibliography}", encoding="utf-8")
            with patch("app.citation_mapper.ainvoke_json", new=AsyncMock(return_value={})), \
                 patch("app.workflow.check_llm", new=AsyncMock(return_value={"status": "ok"})), \
                 patch("app.workflow.ainvoke_json", new=AsyncMock(return_value=review)):
                events = [event async for event in service.stream(path)]
        starts = [event["step"] for event in events if event["type"] == "step_start"]
        ends = [event["step"] for event in events if event["type"] == "step_end"]
        self.assertEqual(starts, ends)
        self.assertEqual(starts[:3], ["agent6_plan", "agent1_extract", "agent2_map"])
        final = [event["report"] for event in events if event["type"] == "final"]
        self.assertEqual(len(final), 1)
        self.assertEqual(final[0]["statistics"]["passed"], 1)

    async def test_failed_review_retry_only_reruns_failed_pair(self):
        service = WorkflowService()
        pairs = [
            {"pair_id": "PAIR00001", "reference_id": "REF0001", "initial_review": {
                "status": "completed", "failure_reason": ""}},
            {"pair_id": "PAIR00002", "reference_id": "REF0002", "initial_review": {
                "status": "error", "failure_reason": "model_error"}},
        ]
        recovered = {"status": "completed", "failure_reason": "", "confidence": .9,
                     "verdict": "supports", "rationale": "recovered"}
        with patch.object(service, "_review", new=AsyncMock(return_value=recovered)) as reviewer, \
             patch("app.workflow.asyncio.sleep", new=AsyncMock()):
            result, retry_log = await service._retry_failed_reviews(
                pairs, {"REF0001": {}, "REF0002": {}}, "initial_review", deep=False)
        reviewer.assert_awaited_once()
        self.assertEqual(reviewer.await_args.args[0]["pair_id"], "PAIR00002")
        self.assertEqual(result[0]["initial_review"]["status"], "completed")
        self.assertEqual(result[1]["initial_review"]["status"], "completed")
        self.assertEqual(len(retry_log), 1)

    async def test_equality_enters_deep_and_can_pass(self):
        report, service, calls = await self.run_case(initial=.75, deep=.9)
        self.assertEqual(report["pairs"][0]["final_decision"]["passed_by"], "deep_review")
        self.assertEqual(calls, 2)

    async def test_transport_failure_is_manual_not_negative(self):
        report, service, calls = await self.run_case(tool_failure=True)
        self.assertEqual(calls, 0)
        self.assertEqual(service.literature.call.await_count, 0)
        self.assertIsNone(report["pairs"][0]["initial_review"]["confidence"])
        self.assertIsNone(report["pairs"][0]["final_decision"]["effective_confidence"])
        self.assertEqual(report["statistics"]["manual_required"], 1)
        self.assertEqual(report["statistics"]["rejected"], 0)
        self.assertEqual(report["status"], "failed")
        self.assertEqual(report["statistics"]["failure_reasons"], {"literature_unavailable": 1})

    async def test_api_preflight_error_is_actionable(self):
        from fastapi.testclient import TestClient
        from app.main import app
        with patch("app.main.check_llm", new=AsyncMock(side_effect=LLMServiceError("connection_error", "Check LLM_BASE_URL"))):
            with TestClient(app) as client:
                response = client.post("/api/check-model")
        self.assertEqual(response.status_code, 503)
        self.assertIn("LLM_BASE_URL", response.json()["detail"])

    async def test_no_abstract_can_recover_from_fulltext(self):
        report, _, calls = await self.run_case(metadata={"found": True, "abstract": "", "access_status": "open-access"})
        self.assertEqual(calls, 1)
        self.assertTrue(report["pairs"][0]["final_decision"]["passed"])

    async def test_closed_source_no_fulltext_is_manual(self):
        report, _, _ = await self.run_case(initial=.4, full_available=False,
            metadata={"found": True, "abstract": "Evidence supports the claim.", "access_status": "closed-source"})
        pair = report["pairs"][0]
        self.assertEqual(pair["standardized_response"]["access_status"], "close-source")
        self.assertEqual(pair["final_decision"]["reason_code"], "full_text_unavailable")
        self.assertIsNone(pair["final_decision"]["effective_confidence"])

    async def test_evidence_rejection_differs_from_model_failure(self):
        report, _, _ = await self.run_case(initial=.3, deep=.2)
        self.assertEqual(report["statistics"]["rejected"], 1)
        for kwargs in ({"llm_failure": True}, {"invalid": True}):
            report, service, _ = await self.run_case(**kwargs)
            self.assertEqual(report["statistics"]["manual_required"], 1)
            self.assertEqual(service.literature.call.await_count, 1)

    def test_scores_are_not_silently_clamped(self):
        for value in (float("nan"), 1.5, -1):
            with self.assertRaises(ValidationError):
                ReviewPayload(confidence=value, verdict="supports", rationale="test")


if __name__ == "__main__":
    unittest.main()
